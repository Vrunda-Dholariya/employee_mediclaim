import frappe
from frappe import _
from frappe.utils import date_diff, nowdate, getdate
from frappe.utils.xlsxutils import make_xlsx


def execute(filters=None):
    filters = frappe._dict(filters or {})
    data = get_data(filters)
    for row in data:
        dob = row.get("date_of_birth")
        row["age"] = date_diff(getdate(nowdate()), getdate(dob)) // 365 if dob else None
        row["member_type"] = "Employee" if (row.get("relation") or "").strip().lower() == "self" else "Dependent"
    return get_columns(), data


def get_columns():
    return [
        {"label": _("Employee Code"), "fieldname": "employee_code", "fieldtype": "Data", "width": 210},
        {"label": _("Employee Name"), "fieldname": "employee_name", "fieldtype": "Data", "width": 260},
        {"label": _("Corporate Name"), "fieldname": "corporate_name", "fieldtype": "Data", "width": 250},
        {"label": _("Policy No"), "fieldname": "policy_no", "fieldtype": "Data", "width": 250},
        {"label": _("Policy Effective Date"), "fieldname": "policy_effective_date", "fieldtype": "Date", "width": 230},
        {"label": _("Company Date of Joining"), "fieldname": "company_doj", "fieldtype": "Date", "width": 230},
        {"label": _("Dependant Name"), "fieldname": "dependant_name", "fieldtype": "Data", "width": 260},
        {"label": _("Relation"), "fieldname": "relation", "fieldtype": "Data", "width": 190},
        {"label": _("Gender"), "fieldname": "gender", "fieldtype": "Link", "options": "Gender", "width": 190},
        {"label": _("Date of Birth"), "fieldname": "date_of_birth", "fieldtype": "Date", "width": 200},
        {"label": _("Age (Years)"), "fieldname": "age", "fieldtype": "Int", "width": 190},
        {"label": _("Sum Insured"), "fieldname": "sum_insured", "fieldtype": "Data", "width": 200},
        {"label": _("Mobile No."), "fieldname": "mobile_no", "fieldtype": "Data", "width": 210},
        {"label": _("Employee's E-mail Id"), "fieldname": "email", "fieldtype": "Data", "width": 260},
        {"label": _("Location"), "fieldname": "location", "fieldtype": "Data", "width": 200},
        {"label": _("Member Type"), "fieldname": "member_type", "fieldtype": "Data", "width": 200},
        {"label": _("Grade"), "fieldname": "grade", "fieldtype": "Link", "options": "Employee Grade", "width": 150},
        {"label": _("Designation"), "fieldname": "designation", "fieldtype": "Link", "options": "Designation", "width": 220},
        {"label": _("Salary Account No"), "fieldname": "salary_account_no", "fieldtype": "Data", "width": 220},
        {"label": _("MICR"), "fieldname": "micr_code", "fieldtype": "Data", "width": 190},
        {"label": _("IFSC"), "fieldname": "ifsc_code", "fieldtype": "Data", "width": 190},
        {"label": _("Bank Name"), "fieldname": "bank_name", "fieldtype": "Data", "width": 200},
        {"label": _("Account Type"), "fieldname": "account_type", "fieldtype": "Data", "width": 190},
        {"label": _("Bank Address"), "fieldname": "bank_address", "fieldtype": "Data", "width": 200},
        {"label": _("Bank Contact No"), "fieldname": "bank_contact_no", "fieldtype": "Data", "width": 190},
        {"label": _("Name as Per Bank Account"), "fieldname": "name_as_per_bank", "fieldtype": "Data", "width": 200},
        {"label": _("Member Effective Date"), "fieldname": "member_effective_date", "fieldtype": "Date", "width": 210},
        {"label": _("Deletion Date"), "fieldname": "deletion_date", "fieldtype": "Date", "width": 200},
    ]


def get_data(filters):
    conditions, values = [], {}

    if filters.get("dependent_addition"):
        conditions.append("mm.relation != 'Self'")
        date_field = "e.date_of_joining"
    elif filters.get("member_deletion"):
        conditions.append("e.status = 'Left'")
        date_field = "e.relieving_date"
    else:
        conditions.append("mm.relation = 'Self'")
        date_field = "e.date_of_joining"

    if filters.get("from_date"):
        conditions.append(f"{date_field} >= %(from_date)s")
        values["from_date"] = filters["from_date"]
    if filters.get("to_date"):
        conditions.append(f"{date_field} <= %(to_date)s")
        values["to_date"] = filters["to_date"]

    if filters.get("employee_filter"):
        conditions.append("e.name = %(employee_filter)s")
        values["employee_filter"] = filters["employee_filter"]

    where = ("WHERE " + " AND ".join(conditions)) if conditions else ""

    return frappe.db.sql(f"""
        SELECT COALESCE(NULLIF(TRIM(e.employee_number), ''), e.name) AS employee_code,
               e.employee_name, e.company AS corporate_name, e.custom_policy_no AS policy_no,
               e.custom_policy_effective_date AS policy_effective_date,
               e.date_of_joining AS company_doj,
               e.custom_date_of_internship AS date_of_internship,
               mm.member_name AS dependant_name, mm.relation, mm.gender, mm.date_of_birth,
               mm.effective_date AS member_effective_date,
               e.relieving_date AS deletion_date,
               e.cell_number AS mobile_no, e.company_email AS email,
               e.grade, e.designation, e.bank_ac_no AS salary_account_no,
               e.micr_code, e.ifsc_code, e.bank_name
        FROM `tabEmployee` e
        JOIN `tabMediclaim Member` mm ON mm.parent = e.name AND mm.parenttype = 'Employee'
        {where}
        ORDER BY e.employee_name, mm.relation
    """, values, as_dict=True)


@frappe.whitelist()
def export_all_views(filters=None):
    import json
    if isinstance(filters, str):
        filters = json.loads(filters)
    filters = frappe._dict(filters or {})

    columns = get_columns()
    header = ["View"] + [c["label"] for c in columns]
    rows = [header]

    views = {
        "employee_addition": "Employee Addition",
        "dependent_addition": "Dependent Addition",
        "member_deletion": "Member Deletion",
    }

    for view_key, view_label in views.items():
        view_filters = frappe._dict(filters.copy())
        view_filters["employee_addition"] = 1 if view_key == "employee_addition" else 0
        view_filters["dependent_addition"] = 1 if view_key == "dependent_addition" else 0
        view_filters["member_deletion"] = 1 if view_key == "member_deletion" else 0

        data = get_data(view_filters)
        for row in data:
            dob = row.get("date_of_birth")
            row["age"] = date_diff(getdate(nowdate()), getdate(dob)) // 365 if dob else None
            row["member_type"] = "Employee" if (row.get("relation") or "").strip().lower() == "self" else "Dependent"
            rows.append([view_label] + [row.get(c["fieldname"]) for c in columns])

        xlsx_file = make_xlsx(rows, "Employee Mediclaim Report")

    from openpyxl import load_workbook
    wb = load_workbook(xlsx_file)
    ws = wb.active

    ws.column_dimensions["A"].width = 18
    date_columns = []  # Excel column letters that hold Date fields

    for i, col in enumerate(columns):
        excel_col_letter = ws.cell(row=1, column=i + 2).column_letter
        pixel_width = col.get("width", 150)
        ws.column_dimensions[excel_col_letter].width = pixel_width / 7
        if col.get("fieldtype") == "Date":
            date_columns.append(excel_col_letter)

    for col_letter in date_columns:
        for row_num in range(2, ws.max_row + 1):
            cell = ws[f"{col_letter}{row_num}"]
            if cell.value:
                cell.number_format = "DD-MM-YYYY"

    import io
    output = io.BytesIO()
    wb.save(output)

    frappe.response["filename"] = "Employee_Mediclaim_Report_All_Views.xlsx"
    frappe.response["filecontent"] = output.getvalue()
    frappe.response["type"] = "binary"